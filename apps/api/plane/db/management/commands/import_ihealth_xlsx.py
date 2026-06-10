# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

import html
import re
from datetime import datetime

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from plane.db.models import (
    DEFAULT_STATES,
    Issue,
    IssueAssignee,
    Project,
    ProjectMember,
    State,
    User,
    Workspace,
    WorkspaceMember,
)

SKIP_SHEETS = {"Weekly(short)", "IT project", "AI_Weekly_Report"}
NAME_HEADERS = {"что сделано", "что сделал"}
STATUS_MAP = {
    "done": "completed",
    "in progress": "started",
    "backlog": "backlog",
}

WORKSPACE_NAME = "iHealth"
WORKSPACE_SLUG = "ihealth"

MEMBER_EMAILS = [
    "admin@pro-ecta.kz",
    "syrym@galamat.com",
    "zhs@galamat.com",
    "dimash@galamat.com",
]
OWNER_EMAIL = "syrym@galamat.com"
ASSIGNEE_EMAIL = "dimash@galamat.com"
ADMIN_EMAILS = {"syrym@galamat.com", "zhs@galamat.com", "admin@pro-ecta.kz"}


def make_identifier(name: str, used: set[str]) -> str:
    translit = {
        "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e",
        "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m",
        "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u",
        "ф": "f", "х": "h", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "sch",
        "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
    }
    lower = name.lower()
    chars = []
    for ch in lower:
        if ch in translit:
            chars.append(translit[ch])
        elif ch.isalnum():
            chars.append(ch)
    base = "".join(chars).upper()[:12] or "PRJ"
    candidate = base
    n = 1
    while candidate in used:
        suffix = str(n)
        candidate = (base[: 12 - len(suffix)] + suffix)[:12]
        n += 1
    used.add(candidate)
    return candidate


def parse_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(value[:19], fmt).date()
            except ValueError:
                continue
    return None


def normalize_status(raw) -> str:
    if raw is None:
        return "unstarted"
    key = str(raw).strip().lower()
    return STATUS_MAP.get(key, "unstarted")


def build_description(insight, next_step) -> str:
    parts = []
    if insight and str(insight).strip():
        parts.append(f"<p><strong>Инсайт / результат:</strong> {html.escape(str(insight).strip())}</p>")
    if next_step and str(next_step).strip():
        parts.append(f"<p><strong>Следующий микро-шаг:</strong> {html.escape(str(next_step).strip())}</p>")
    return "".join(parts) if parts else ""


def is_header_row(row) -> bool:
    if not row:
        return False
    first = str(row[0] or "").strip().lower()
    return first in ("дата", "date")


def should_skip_sheet(name: str) -> bool:
    if name in SKIP_SHEETS:
        return True
    if "roadmap" in name.lower():
        return True
    return False


class Command(BaseCommand):
    help = "Import iHealth workspace, projects and issues from Excel progress file"

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", type=str, help="Path to Excel file")
        parser.add_argument("--owner", type=str, default=OWNER_EMAIL, help="Workspace owner email")
        parser.add_argument("--force", action="store_true", help="Re-import issues into existing workspace")

    def handle(self, *args, **options):
        xlsx_path = options["xlsx_path"]
        owner_email = options["owner"]
        force = options["force"]

        owner = User.objects.filter(email=owner_email).first()
        if owner is None:
            raise CommandError(f"Owner user {owner_email} not found. Run setup_instance_users first.")

        members = list(User.objects.filter(email__in=MEMBER_EMAILS))
        assignee = User.objects.filter(email=ASSIGNEE_EMAIL).first()
        if assignee is None:
            raise CommandError(f"Assignee {ASSIGNEE_EMAIL} not found.")

        try:
            wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        except Exception as e:
            raise CommandError(f"Cannot open {xlsx_path}: {e}") from e

        with transaction.atomic():
            workspace = Workspace.objects.filter(slug=WORKSPACE_SLUG).first()
            if workspace and not force:
                self.stdout.write(self.style.WARNING(f"Workspace {WORKSPACE_SLUG} exists. Use --force to import issues."))
                return

            if workspace is None:
                workspace = Workspace.objects.create(
                    name=WORKSPACE_NAME,
                    slug=WORKSPACE_SLUG,
                    owner=owner,
                    created_by=owner,
                )
                self.stdout.write(self.style.SUCCESS(f"Created workspace {WORKSPACE_NAME}"))

            for user in members:
                role = 20 if user.email in ADMIN_EMAILS else 15
                WorkspaceMember.objects.get_or_create(
                    workspace=workspace,
                    member=user,
                    defaults={"role": role, "created_by": owner},
                )

            used_identifiers: set[str] = set(
                Project.objects.filter(workspace=workspace).values_list("identifier", flat=True)
            )
            total_issues = 0

            for sheet_name in wb.sheetnames:
                if should_skip_sheet(sheet_name):
                    continue

                ws = wb[sheet_name]
                project = Project.objects.filter(workspace=workspace, name=sheet_name).first()
                if project is None:
                    identifier = make_identifier(sheet_name, used_identifiers)
                    project = Project.objects.create(
                        workspace=workspace,
                        name=sheet_name,
                        identifier=identifier,
                        created_by=owner,
                    )
                    State.objects.bulk_create(
                        [
                            State(
                                name=state["name"],
                                color=state["color"],
                                project=project,
                                sequence=state["sequence"],
                                workspace=workspace,
                                group=state["group"],
                                default=state.get("default", False),
                                created_by=owner,
                            )
                            for state in DEFAULT_STATES
                        ]
                    )
                    for user in members:
                        role = 20 if user.email in ADMIN_EMAILS else 15
                        ProjectMember.objects.get_or_create(
                            project=project,
                            member=user,
                            defaults={"workspace": workspace, "role": role, "created_by": owner},
                        )
                    self.stdout.write(self.style.SUCCESS(f"Created project {sheet_name} ({identifier})"))
                elif force:
                    Issue.objects.filter(project=project).delete()

                states_by_group = {s.group: s for s in State.objects.filter(project=project)}

                row_count = 0
                for idx, row in enumerate(ws.iter_rows(values_only=True)):
                    if idx == 0 or is_header_row(row):
                        continue
                    if not row or len(row) < 2:
                        continue
                    task_name = row[1]
                    if task_name is None or not str(task_name).strip():
                        continue
                    if str(task_name).strip().lower() in NAME_HEADERS:
                        continue

                    task_date = parse_date(row[0])
                    insight = row[2] if len(row) > 2 else None
                    next_step = row[3] if len(row) > 3 else None
                    status_raw = row[4] if len(row) > 4 else None
                    group = normalize_status(status_raw)
                    state = states_by_group.get(group) or states_by_group.get("unstarted")

                    issue = Issue(
                        project=project,
                        workspace=workspace,
                        name=str(task_name).strip()[:255],
                        description_html=build_description(insight, next_step),
                        state=state,
                        target_date=task_date,
                        created_by=owner,
                        updated_by=owner,
                    )
                    issue.save()
                    if group == "completed" and task_date:
                        Issue.objects.filter(pk=issue.pk).update(
                            completed_at=timezone.make_aware(
                                datetime.combine(task_date, datetime.min.time())
                            )
                        )

                    IssueAssignee.objects.get_or_create(
                        issue=issue,
                        assignee=assignee,
                        project=project,
                        workspace=workspace,
                        defaults={"created_by": owner, "updated_by": owner},
                    )
                    row_count += 1

                total_issues += row_count
                self.stdout.write(f"  {sheet_name}: {row_count} issues")

        self.stdout.write(self.style.SUCCESS(f"Import complete. Total issues: {total_issues}"))
